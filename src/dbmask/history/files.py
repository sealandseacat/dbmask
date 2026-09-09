"""CSV, XLSX and a strict single-table Markdown interchange format."""
from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from dbmask.history.records import (
    HISTORY_FIELDS,
    REQUIRED_HEADERS,
    HistoryRecord,
    HistoryValidationError,
)


def _excel():
    try:
        import openpyxl
    except ImportError as exc:
        raise HistoryValidationError('XLSX support requires: pip install "dbmask[excel]"') from exc
    return openpyxl


def _markdown_cells(line: str) -> list[str]:
    if not line.startswith("|") or not line.endswith("|"):
        raise HistoryValidationError("Markdown must contain exactly one pipe-delimited table")
    cells = re.split(r"(?<!\\)\|", line[1:-1])
    return [c.strip().replace(r"\|", "|").replace("<br>", "\n") for c in cells]


def read_history_file(path: str | Path, *, sheet: str = "history") -> list[HistoryRecord]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, strict=True))
        except csv.Error as exc:
            raise HistoryValidationError(f"Invalid CSV: {exc}") from exc
    elif suffix == ".md":
        lines = [line.strip() for line in path.read_text(encoding="utf-8-sig").splitlines()
                 if line.strip()]
        if len(lines) < 2:
            raise HistoryValidationError("Markdown requires a header and separator row")
        rows = [_markdown_cells(line) for line in lines]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
            raise HistoryValidationError("Invalid Markdown separator row")
        if len(rows[1]) != len(rows[0]):
            raise HistoryValidationError("Markdown separator width does not match header")
        del rows[1]
    elif suffix == ".xlsx":
        book = _excel().load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet not in book.sheetnames:
                raise HistoryValidationError(f"Missing Excel worksheet {sheet!r}")
            rows = []
            for row in book[sheet].iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise HistoryValidationError("Excel formulas are not allowed in history records")
                rows.append([cell.value for cell in row])
        finally:
            book.close()
    else:
        raise HistoryValidationError("Supported history formats: .csv, .xlsx, .md")
    if not rows:
        raise HistoryValidationError("History file is empty")
    headers = [str(value or "").strip() for value in rows[0]]
    # Excel formatting can create trailing empty columns; ignore only those.
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) != len(set(headers)) or any(not h for h in headers):
        raise HistoryValidationError("History headers must be unique and nonempty")
    missing = set(REQUIRED_HEADERS) - set(headers)
    unknown = set(headers) - set(HISTORY_FIELDS)
    if missing or unknown:
        raise HistoryValidationError(f"Invalid headers: missing={sorted(missing)}, unknown={sorted(unknown)}")
    records = []
    for line_number, values in enumerate(rows[1:], 2):
        if all(value is None or value == "" for value in values):
            continue
        if len(values) < len(headers) or any(v not in (None, "") for v in values[len(headers):]):
            raise HistoryValidationError(f"Row {line_number}: wrong number of cells")
        data: dict[str, Any] = {}
        for key, value in zip(headers, values):
            if key == "revision":
                raw = str(value) if value not in (None, "") else "0"
                if not re.fullmatch(r"\d+", raw):
                    raise HistoryValidationError(f"Row {line_number}: revision must be an integer")
                data[key] = int(raw)
            elif isinstance(value, (datetime, date)) and key in {"analysis_date", "reviewed_at", "expires_at"}:
                data[key] = value.isoformat()
            elif value is None:
                data[key] = ""
            elif not isinstance(value, str):
                raise HistoryValidationError(
                    f"Row {line_number}: {key} must be text (format IDs as Text in Excel)"
                )
            else:
                data[key] = value if key in {"database", "schema", "table", "column", "user_id"} else value.strip()
        try:
            records.append(HistoryRecord(**data))
        except TypeError as exc:
            raise HistoryValidationError(f"Row {line_number}: {exc}") from exc
    return records


def write_history_file(
    path: str | Path, records: list[HistoryRecord], *, overwrite: bool = False,
) -> None:
    """Write literal strings, including formula-looking text, without evaluation."""
    path = Path(path)
    if path.exists() and not overwrite:
        raise HistoryValidationError(f"Output already exists: {path}")
    rows = [[record.to_dict()[field] for field in HISTORY_FIELDS] for record in records]
    suffix = path.suffix.lower()
    if suffix == ".csv":
        # CSV cannot encode a cell's text type. Reject unsafe spreadsheet prefixes
        # instead of modifying an identifier, method or reason during a round trip.
        if any(isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@"))
               for row in rows for v in row):
            raise HistoryValidationError("Formula-looking text cannot be safely exported to CSV; use XLSX")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(HISTORY_FIELDS)
            writer.writerows(rows)
    elif suffix == ".md":
        def cell(value):
            text = str(value)
            if "\\" in text or "<br>" in text or text != text.strip():
                raise HistoryValidationError("Markdown cannot round-trip this text; use XLSX or CSV")
            return text.replace("|", r"\|").replace("\n", "<br>")
        lines = ["| " + " | ".join(HISTORY_FIELDS) + " |",
                 "| " + " | ".join("---" for _ in HISTORY_FIELDS) + " |"]
        lines.extend("| " + " | ".join(cell(v) for v in row) + " |" for row in rows)
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    elif suffix == ".xlsx":
        excel = _excel()
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        book = excel.Workbook()
        ws = book.active
        ws.title = "history"
        for row in [list(HISTORY_FIELDS), *rows]:
            ws.append(row)
        for row in ws.iter_rows():
            for item in row:
                if isinstance(item.value, str):
                    item.data_type = "s"
                    item.number_format = "@"
        for item in ws[1]:
            item.font = Font(bold=True, color="FFFFFF")
            item.fill = PatternFill("solid", fgColor="245B83")
        widths = [22, 18, 24, 24, 14, 20, 24, 18, 20, 20, 28, 55, 28, 22, 28, 12]
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[excel.utils.get_column_letter(index)].width = width
        for column, options in (("E", "mask,keep,review"), ("H", "pending,approved,superseded")):
            validation = DataValidation(type="list", formula1=f'"{options}"')
            validation.errorTitle = "Choose a listed value"
            validation.error = options
            validation.showErrorMessage = True
            ws.add_data_validation(validation)
            validation.add(f"{column}2:{column}10000")
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
        book.save(path)
        book.close()
    else:
        raise HistoryValidationError("Supported history formats: .csv, .xlsx, .md")
