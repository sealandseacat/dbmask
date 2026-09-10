"""Exercise the real review cycle, exact scope and failure-before-replacement."""
from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml

from dbmask.cli import cli
from dbmask.config import Config
from dbmask.detection.overrides import FieldOverrides
from dbmask.history.file_store import FileHistoryStore
from dbmask.history.files import read_history_file, write_history_file
from dbmask.history.records import HistoryConflictError, HistoryRecord, HistoryValidationError
from dbmask.history.writeback import export_review, manifest_path, writeback
from dbmask.llm.base import LLMResult
from dbmask.runner import Runner


def record(**changes):
    return replace(HistoryRecord(
        'sample', 'main', 'customers', 'email', decision='mask', detected_type='email',
        masking_strategy='blank', user_id='00123', analysis_date='2026-09-09', data_type='TEXT',
    ), **changes)


def approved(**changes):
    return record(review_status='approved', reviewed_by='original-reviewer',
                  reviewed_at='2026-09-09T10:00:00Z', **changes)


def review_pair(tmp_path, *, source_rows=None, review_rows=None, suffix='csv'):
    source = tmp_path / f'master.{suffix}'
    output = tmp_path / f'review.{suffix}'
    write_history_file(source, source_rows or [])
    store = FileHistoryStore(source)
    store.connect()
    rows = review_rows if review_rows is not None else [record()]
    export_review(store, output, rows)
    write_history_file(output, [replace(r, review_status='approved') for r in rows], overwrite=True)
    return store, output


@pytest.mark.parametrize('suffix', ['csv', 'xlsx', 'md'])
def test_cli_round_trip_original_file_is_authoritative_and_other_database_unchanged(cli_env, suffix):
    master = cli_env.tmp_path / f'master.{suffix}'
    other = approved(database='database-B', decision='keep', masking_strategy='')
    write_history_file(master, [other])
    before = master.read_bytes()
    database_before = cli_env.db.read_bytes()
    cfg = cli_env.make_config(history={'source_file': master.name}, detection={'user_id': 'analyst-A'})
    result = cli_env.runner.invoke(cli, ['history-import', '--config', cfg, '--file', str(master)])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)['written'] is False
    review = cli_env.tmp_path / 'review.xlsx'
    result = cli_env.runner.invoke(cli, ['scan', '--config', cfg, '--output', str(review)])
    assert result.exit_code == 0, result.output
    assert master.read_bytes() == before
    assert cli_env.db.read_bytes() == database_before
    assert not (cli_env.tmp_path / 'history.db').exists()
    rows = read_history_file(review)
    assert all(r.review_status == 'pending' for r in rows)
    rows = [replace(r, review_status='approved', masking_strategy='blank', reason='Human override')
            if r.column == 'email' else r for r in rows]
    write_history_file(review, rows, overwrite=True)
    args = ['history-writeback', '--config', cfg, '--file', str(review), '--reviewed-by', 'siyuan']
    result = cli_env.runner.invoke(cli, args)
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)['inserted'] == 1
    assert master.read_bytes() == before
    result = cli_env.runner.invoke(cli, [*args, '--apply'])
    assert result.exit_code == 0, result.output
    assert Path(json.loads(result.output)['backup']).read_bytes() == before
    saved = {r.database: r for r in read_history_file(master)}
    assert saved['database-B'] == other
    assert saved['sample'].reviewed_by == 'siyuan'
    assert saved['sample'].user_id == 'analyst-A'
    assert saved['sample'].masking_strategy == 'blank'
    assert saved['sample'].revision == 1
    with Runner(Config.load(cfg)) as runner:
        decision = next(d for d in runner.scan().decisions if d.column == 'email')
        assert decision.source == 'history' and decision.masking_strategy == 'blank'
        assert runner.pipeline.stats.by_source['history'] == 1
    cfg_b = cli_env.make_config(name='b.yaml', database={'name': 'database-B'},
                               history={'source_file': master.name})
    with Runner(Config.load(cfg_b)) as runner:
        decision = next(d for d in runner.scan().decisions if d.column == 'email')
        assert decision.source == 'history' and not decision.is_sensitive
    result = cli_env.runner.invoke(cli, ['mask', '--config', cfg, '--apply'])
    assert result.exit_code == 0, result.output
    assert cli_env.read('email') == ['', '', '']
    assert read_history_file(master) == list(saved.values())


def test_llm_suggestion_is_reviewed_overridden_and_reused_without_llm(cli_env, monkeypatch):
    master = cli_env.tmp_path / 'master.md'
    write_history_file(master, [])
    calls = []

    def classify(column, values):
        calls.append(column)
        return LLMResult(True, 'free_text', .7)

    monkeypatch.setattr('dbmask.runner.create_provider', lambda cfg: SimpleNamespace(classify=classify))
    cfg = cli_env.make_config(history={'source_file': str(master)}, detection={'user_id': 'analyst'})
    review = cli_env.tmp_path / 'review.csv'
    result = cli_env.runner.invoke(cli, ['scan', '--config', cfg, '--output', str(review)])
    assert result.exit_code == 0, result.output
    assert 'status_code' in calls and 'email' not in calls  # patterns ran first
    rows = read_history_file(review)
    status = next(r for r in rows if r.column == 'status_code')
    assert status.decision == 'mask' and status.review_status == 'pending'
    rows = [replace(r, decision='keep', masking_strategy='', review_status='approved',
                    reason='Human confirms internal status') if r == status else r for r in rows]
    write_history_file(review, rows, overwrite=True)
    result = cli_env.runner.invoke(cli, ['history-writeback', '--config', cfg, '--file', str(review),
                                        '--reviewed-by', 'human', '--apply'])
    assert result.exit_code == 0, result.output
    calls.clear()
    with Runner(Config.load(cfg)) as runner:
        result = next(d for d in runner.scan().decisions if d.column == 'status_code')
        assert result.source == 'history' and not result.is_sensitive
    assert 'status_code' not in calls


@pytest.mark.parametrize('change', [{'database': 'Sample'}, {'schema': 'Main'},
                                  {'table': 'Customers'}, {'column': 'Email'},
                                  {'database': 'sample.main', 'schema': ''}])
def test_file_lookup_uses_case_preserving_full_tuple(tmp_path, change):
    store, _ = review_pair(tmp_path, source_rows=[approved()])
    assert store.get('sample', 'main', 'customers', 'email', current_type='text').source == 'history'
    other = replace(approved(), **change)
    assert store.get(other.database, other.schema, other.table, other.column, current_type='TEXT') is None


@pytest.mark.parametrize('change', [{'decision': 'keep', 'masking_strategy': ''},
                                  {'masking_strategy': 'redact'}, {'reason': 'Updated rationale'}])
def test_approved_replacement_has_revision_and_new_reviewer(tmp_path, change):
    old = approved(revision=4)
    store, review = review_pair(tmp_path, source_rows=[old], review_rows=[replace(old, **change)])
    result = writeback(store, review, reviewed_by='new-reviewer', apply=True)
    assert result['updated'] == 1
    new = read_history_file(store.path)[0]
    assert new.revision == 5 and new.reviewed_by == 'new-reviewer'
    assert Path(result['backup']).read_bytes() == store.snapshot


def test_fresh_unchanged_approval_is_noop_and_pending_rows_are_not_saved(tmp_path):
    old = approved(revision=2)
    store, review = review_pair(tmp_path, source_rows=[old], review_rows=[old, record(column='city')])
    write_history_file(review, [old, record(column='city')], overwrite=True)
    result = writeback(store, review, reviewed_by='another-reviewer', apply=True)
    assert result['unchanged'] == 1 and result['skipped'] == 1 and result['backup'] is None
    assert store.path.read_bytes() == store.snapshot
    assert not list(tmp_path.glob('*.bak'))


@pytest.mark.parametrize('change', [{'database': 'B'}, {'column': 'different'}, {'revision': 99},
                                  {'user_id': 'different'}, {'analysis_date': '2026-09-10'},
                                  {'data_type': 'VARCHAR(10)'}])
def test_edited_identity_or_scan_baseline_is_rejected_without_partial_changes(tmp_path, change):
    store, review = review_pair(tmp_path, review_rows=[record(), record(column='city')])
    rows = read_history_file(review)
    write_history_file(review, [rows[0], replace(rows[1], **change)], overwrite=True)
    with pytest.raises(HistoryConflictError):
        writeback(store, review, reviewed_by='human', apply=True)
    assert store.path.read_bytes() == store.snapshot
    assert not list(tmp_path.glob('*.bak'))


@pytest.mark.parametrize('change', [{'masking_strategy': 'does_not_exist'}, {'expires_at': '2000-01-01'},
                                  {'decision': 'keep'}, {'decision': 'review'},
                                  {'review_status': 'APPROVED'}])
def test_invalid_approval_rejects_entire_batch(tmp_path, change):
    store, review = review_pair(tmp_path, review_rows=[record(), record(column='city')])
    rows = read_history_file(review)
    write_history_file(review, [rows[0], replace(rows[1], **change)], overwrite=True)
    with pytest.raises(HistoryValidationError):
        writeback(store, review, reviewed_by='human', apply=True)
    assert store.path.read_bytes() == store.snapshot


def test_duplicates_in_master_and_review_are_rejected(tmp_path):
    store, review = review_pair(tmp_path)
    rows = read_history_file(review)
    write_history_file(review, rows * 2, overwrite=True)
    with pytest.raises(HistoryConflictError, match='Duplicate'):
        writeback(store, review, reviewed_by='human', apply=True)
    write_history_file(store.path, [approved(), approved()], overwrite=True)
    with pytest.raises(HistoryConflictError, match='Duplicate'):
        FileHistoryStore(store.path).connect()


@pytest.mark.parametrize('when', ['before_open', 'before_apply'])
def test_changed_master_rejects_stale_review(tmp_path, when):
    store, review = review_pair(tmp_path)
    write_history_file(store.path, [approved(database='someone-else')], overwrite=True)
    changed = store.path.read_bytes()
    if when == 'before_open':
        store.connect()
    with pytest.raises(HistoryConflictError, match='changed'):
        writeback(store, review, reviewed_by='human', apply=True)
    assert store.path.read_bytes() == changed


@pytest.mark.parametrize('failure', ['locked', 'replace', 'backup', 'staging'])
def test_failed_write_keeps_master_and_cleans_own_lock(tmp_path, monkeypatch, failure):
    store, review = review_pair(tmp_path)
    lock = Path(str(store.path) + '.dbmask.lock')

    def fail(*args, **kwargs):
        raise OSError('simulated I/O failure')

    if failure == 'locked':
        lock.write_text('another writer')
    elif failure == 'replace':
        monkeypatch.setattr('dbmask.history.writeback.os.replace', fail)
    elif failure == 'staging':
        monkeypatch.setattr('dbmask.history.writeback.write_history_file', fail)
    else:
        original = Path.open

        def fail_backup(path, *args, **kwargs):
            if path.suffix == '.bak':
                fail()
            return original(path, *args, **kwargs)

        monkeypatch.setattr(Path, 'open', fail_backup)
    with pytest.raises((HistoryConflictError, OSError)):
        writeback(store, review, reviewed_by='human', apply=True)
    assert store.path.read_bytes() == store.snapshot
    assert lock.exists() == (failure == 'locked')
    assert not list(tmp_path.glob('.dbmask-*.csv'))


@pytest.mark.parametrize('problem', ['missing', 'malformed', 'other_source'])
def test_review_companion_is_required_and_bound_to_source(tmp_path, problem):
    store, review = review_pair(tmp_path)
    companion = manifest_path(review)
    if problem == 'missing':
        companion.unlink()
    elif problem == 'malformed':
        companion.write_text('{}')
    else:
        data = json.loads(companion.read_text())
        data['source_file'] = '/another/master.csv'
        companion.write_text(json.dumps(data))
    with pytest.raises(HistoryValidationError):
        writeback(store, review, reviewed_by='human', apply=True)
    assert store.path.read_bytes() == store.snapshot


def test_xlsx_preserves_other_sheets_formulas_styles_and_custom_history_sheet(tmp_path):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    master = tmp_path / 'master.xlsx'
    other = approved(database='B')
    write_history_file(master, [other])
    book = openpyxl.load_workbook(master)
    book['history'].title = 'Decisions'
    book['Decisions']['A2'].font = Font(bold=True, color='FF0000')
    notes = book.create_sheet('Notes')
    notes['A1'] = '=1+1'
    notes['B1'] = 'Keep this worksheet'
    notes['B1'].comment = Comment('Do not drop this', 'siyuan')
    book.save(master)
    book.close()
    store = FileHistoryStore(master, sheet='Decisions')
    store.connect()
    review = tmp_path / 'review.csv'
    export_review(store, review, [record()])
    # Use XLSX for formula-looking literal text, since CSV deliberately rejects it.
    review = tmp_path / 'review.xlsx'
    export_review(store, review, [record()])
    write_history_file(review, [approved(reason='=literal text')], overwrite=True)
    writeback(store, review, reviewed_by='human', apply=True)
    book = openpyxl.load_workbook(master, data_only=False)
    try:
        assert book.sheetnames == ['Decisions', 'Notes']
        assert book['Notes']['A1'].value == '=1+1' and book['Notes']['A1'].data_type == 'f'
        assert book['Notes']['B1'].comment.text == 'Do not drop this'
        assert book['Decisions']['A2'].font.bold
        assert book['Decisions']['L3'].value == '=literal text'
        assert book['Decisions']['L3'].data_type == 's'
    finally:
        book.close()
    assert read_history_file(master, sheet='Decisions')[0] == other


def test_type_drift_requires_review_and_export_captures_current_type(cli_env):
    master = cli_env.tmp_path / 'master.csv'
    write_history_file(master, [approved(data_type='VARCHAR(20)', revision=3)])
    before = master.read_bytes()
    cfg = cli_env.make_config(history={'source_file': str(master)}, detection={'user_id': 'new-analyst'})
    review = cli_env.tmp_path / 'review.csv'
    result = cli_env.runner.invoke(cli, ['scan', '--config', cfg, '--output', str(review)])
    assert result.exit_code == 0, result.output
    rows = read_history_file(review)
    row = next(r for r in rows if r.column == 'email')
    assert row.review_status == 'pending' and row.data_type == 'TEXT'
    assert 'VARCHAR(20) -> TEXT' in row.reason
    assert row.user_id == 'new-analyst' and row.revision == 3
    assert master.read_bytes() == before
    write_history_file(review, [replace(r, review_status='approved') if r == row else r for r in rows], overwrite=True)
    result = cli_env.runner.invoke(cli, ['history-writeback', '--config', cfg, '--file', str(review),
                                        '--reviewed-by', 'human', '--apply'])
    assert result.exit_code == 0, result.output
    with Runner(Config.load(cfg)) as runner:
        decision = next(d for d in runner.scan().decisions if d.column == 'email')
        assert decision.source == 'history' and decision.data_type == 'TEXT'


def test_scoped_override_is_exported_and_replaces_approval_only_after_review(cli_env):
    master = cli_env.tmp_path / 'master.csv'
    write_history_file(master, [approved(decision='keep', masking_strategy='', revision=2)])
    overrides = cli_env.tmp_path / 'overrides.yaml'
    overrides.write_text(yaml.safe_dump({'sensitive': [{
        'database': 'sample', 'schema': 'main', 'table': 'customers', 'column': 'email',
        'rule': 'email', 'masking_strategy': 'redact', 'note': 'Override old keep decision',
    }]}))
    cfg = cli_env.make_config(history={'source_file': str(master)},
                              detection={'user_id': 'analyst', 'overrides_file': str(overrides)})
    review = cli_env.tmp_path / 'review.csv'
    result = cli_env.runner.invoke(cli, ['scan', '--config', cfg, '--output', str(review)])
    assert result.exit_code == 0, result.output
    rows = read_history_file(review)
    row = next(r for r in rows if r.column == 'email')
    assert row.decision == 'mask' and row.masking_strategy == 'redact'
    assert row.revision == 2 and row.review_status == 'pending'
    assert read_history_file(master)[0].decision == 'keep'
    write_history_file(review, [replace(r, review_status='approved') if r == row else r for r in rows], overwrite=True)
    result = cli_env.runner.invoke(cli, ['history-writeback', '--config', cfg, '--file', str(review),
                                        '--reviewed-by', 'human', '--apply'])
    assert result.exit_code == 0, result.output
    config = Config.load(cfg)
    config.detection.overrides_file = None
    with Runner(config) as runner:
        row = next(d for d in runner.scan().decisions if d.column == 'email')
        assert row.source == 'history' and row.masking_strategy == 'redact'
    assert read_history_file(master)[0].revision == 3
    obj = FieldOverrides.load(overrides)
    assert obj.decide('B', 'main', 'customers', 'email') is None
    assert obj.decide('Sample', 'main', 'customers', 'email') is None


@pytest.mark.parametrize('raw', [
    {'database': 'A', 'column': 'email'},
    {'database': 'A', 'schema': '', 'table': 't', 'column': 'c', 'match': '*'},
    {'database': 123, 'schema': '', 'table': 't', 'column': 'c'},
    {'match': 'email', 'masking_strategy': 'unknown-strategy'},
])
def test_bad_scoped_overrides_are_not_silently_broadened(raw):
    with pytest.raises(ValueError):
        FieldOverrides()._ingest({'sensitive': [raw]})


def test_scoped_dots_do_not_alias_and_duplicates_are_rejected():
    obj = FieldOverrides()
    raw = {'database': 'a.b', 'schema': 'c', 'table': 't', 'column': 'email'}
    obj._ingest({'sensitive': [raw]})
    assert obj.decide('a.b', 'c', 't', 'email').is_sensitive
    assert obj.decide('a', 'b.c', 't', 'email') is None
    with pytest.raises(ValueError, match='Duplicate'):
        obj._ingest({'not_sensitive': [raw]})


def test_export_never_overwrites_master_or_review_companion(tmp_path):
    store, review = review_pair(tmp_path)
    for target in (store.path, review):
        with pytest.raises(HistoryValidationError):
            export_review(store, target, [record()])
    review.unlink()
    with pytest.raises(HistoryValidationError):
        export_review(store, review, [record()])
    assert store.path.read_bytes() == store.snapshot


def test_file_cli_errors_and_export_without_source_database(cli_env):
    master = cli_env.tmp_path / 'master.csv'
    write_history_file(master, [approved()])
    cfg = cli_env.make_config(history={'source_file': str(master)})
    review = cli_env.tmp_path / 'review.csv'
    missing_user = cli_env.runner.invoke(cli, ['scan', '--config', cfg, '--output', str(review)])
    assert missing_user.exit_code != 0 and 'detection.user_id' in missing_user.output
    result = cli_env.runner.invoke(cli, ['history-export', '--config', cfg, '--output', str(review)])
    assert result.exit_code == 0 and manifest_path(review).exists()
    imported = cli_env.runner.invoke(cli, ['history-import', '--config', cfg, '--file', str(review)])
    assert imported.exit_code != 0 and 'history-writeback' in imported.output
    audited = cli_env.runner.invoke(cli, ['history', '--config', cfg, '--audit'])
    assert audited.exit_code != 0 and 'SQL-only' in audited.output
    cfg_sql = cli_env.make_config(name='sql.yaml')
    result = cli_env.runner.invoke(cli, ['history-writeback', '--config', cfg_sql, '--file', str(review),
                                        '--reviewed-by', 'human'])
    assert result.exit_code != 0 and 'source_file' in result.output
